from read_holdings_excel import get_tlv_data
from pandas import DataFrame
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows  
from openpyxl import load_workbook  
from openpyxl.styles import PatternFill

def column_formula(ws,formula:str, col:int, start_row:int, end_row:int):
    """Apply a formula to a column in an Excel worksheet."""
    for row in range(start_row, end_row + 1):
        cell = ws.cell(row=row, column=col)
        cell.value = formula.replace("##", str(row))

def count_tlv_rows(ws):
    num_rows = 0
    cell = ""
    for row in range(2, ws.max_row):
        cell = str(ws.cell(row=row, column=2).value).strip()
        print(f"Row {row}, Column 2: {cell}")
        if(cell.isnumeric()):
            num_rows += 1
        else:
            break
    return num_rows


def add_df_columns_to_excel(ws, df,df_heading,ws_column):
    """Add columns from a DataFrame to an Excel worksheet."""
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    for idx, col in enumerate(df.columns):
        if col in df_heading:
            col_index = ws_column + df_heading.index(col)
            for row_idx, value in enumerate(df[col], start=2):
                ws.cell(row=row_idx, column=col_index, value=value)
                ws.cell(row=row_idx, column=col_index).fill = yellow_fill  # Highlight the cell with yellow color


def make_performance_xl(file_path: str, output_path: str) -> None:
    """Read holdings from an Excel file, enrich them with Bizportal data, and save to a new Excel file."""
    holdings = get_tlv_data(file_path)
    #holdings = DataFrame()
    # read an excel file from the output path
    workbook = load_workbook(output_path)  
    # get active sheet
    ws = workbook.active
    #print(holdings.columns)
    # 'מספר נייר', 'שם הנייר', 'שווי אחזקה ב ₪', '%  מהתיק', 'פרופיל חשיפה',
    #   'תאריך פעולה אחרונה', '% שינוי יומי', 'תשואה 12 חודשים', 'רווח ב-%',
    #   'רווח ב ₪', 'שער בסיס', 'מס התיק', 'fund_name', 'management_fee',
    #   'three_year_return_percent', 'sharpe_ratio_12_months', 'bizportal_url'
    df = DataFrame(holdings, columns=['שם הנייר', 'מספר נייר', 'three_year_return_percent', 'תשואה 12 חודשים', 'פרופיל חשיפה',
                                      'sharpe_ratio_12_months','management_fee'])

    adjust_rows(ws, df)
    add_df_columns_to_excel(ws, df,'שם הנייר', 1)
    add_df_columns_to_excel(ws, df,'מספר נייר', 2)
    add_df_columns_to_excel(ws, df,'three_year_return_percent', 5)
    add_df_columns_to_excel(ws, df,'תשואה 12 חודשים', 6)
    add_df_columns_to_excel(ws, df,'פרופיל חשיפה', 7)
    add_df_columns_to_excel(ws, df,'sharpe_ratio_12_months', 8)
    add_df_columns_to_excel(ws, df,'management_fee', 9)
    add_df_columns_to_excel(ws, df,'שווי אחזקה ב ₪', 13)
    new_out = output_path.replace(".xlsx", "_MAGIC.xlsx")
    
    workbook.save(new_out)


def adjust_rows(ws, df):
    count = count_tlv_rows(ws)
    print(f"Number of rows with numeric values in column 2: {count}")
    df_row_count = len(df)
    print(f"Number of rows in the DataFrame: {df_row_count}")
    if(df_row_count < count):
        # need to delete the last tlv rows from the worksheet    
        for row in range(df_row_count+2, df_row_count - (count - df_row_count)+2, -1):
                for col in (1,2,5,6,7,8,9,13):
                    ws.cell(row=row, column=col).value = ""

                print(f"cleared row {row} from the worksheet.")
    elif(df_row_count > count):
        # need to add new rows to the worksheet
        for row in range(count+2, df_row_count+2):
            ws.insert_rows(row)
            print(f"Inserted row {row} into the worksheet.")
        
    # 5. Save the Excel file
    #workbook.save(output_file_path)

if __name__ == "__main__":
    input_file_path = "C:/Users/danys/OneDrive/Documents/scripts/finance_data/אחזקות.xls"
    output_file_path = "C:/Users/danys/OneDrive/Documents/scripts/finance_data/דוגמת תחבצ לדני_2 .xlsx"
    make_performance_xl(input_file_path, output_file_path)    


